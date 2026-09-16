"""Generate the .pptx and .docx fixtures.

Content is deliberately distinct per slide/section so tests can assert that a
given phrase resolves to a specific source -- the citation guarantee.
"""

import pathlib

from docx import Document as Docx
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = pathlib.Path(__file__).parent

SLIDES = [
    (
        "Cell Structure",
        ["The cell is the basic unit of life.", "Prokaryotes lack a nucleus."],
        "Remind students that ribosomes are present in both cell types.",
    ),
    (
        "The Mitochondrion",
        ["Mitochondria generate ATP.", "They have a double membrane."],
        "The cristae increase surface area for the electron transport chain.",
    ),
    (
        "Cell Division",
        ["Mitosis produces two identical daughter cells."],
        "",
    ),
]


def build_pptx(path: pathlib.Path) -> int:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # Title and Content

    for title, bullets, notes in SLIDES:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title

        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for extra in bullets[1:]:
            paragraph = body.add_paragraph()
            paragraph.text = extra

        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    # A table on the last slide, to prove table text is captured.
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Comparison"
    table = slide.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(6), Inches(1.2)
    ).table
    table.cell(0, 0).text = "Feature"
    table.cell(0, 1).text = "Prokaryote"
    table.cell(1, 0).text = "Nucleus"
    table.cell(1, 1).text = "Absent"

    presentation.save(path)
    return path.stat().st_size


DOC_SECTIONS = [
    ("Introduction to Genetics", ["Genetics studies heredity and variation."]),
    ("Mendelian Inheritance", [
        "Mendel described dominant and recessive alleles.",
        "A Punnett square predicts offspring genotypes.",
    ]),
    ("Molecular Genetics", ["DNA polymerase replicates the DNA strand."]),
]


def build_docx(path: pathlib.Path) -> int:
    document = Docx()
    document.add_heading("Genetics Handout", level=0)

    for heading, paragraphs in DOC_SECTIONS:
        document.add_heading(heading, level=1)
        for text in paragraphs:
            document.add_paragraph(text)

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Allele"
    table.cell(0, 1).text = "Effect"
    table.cell(1, 0).text = "Recessive"
    table.cell(1, 1).text = "Masked by dominant"

    document.save(path)
    return path.stat().st_size


if __name__ == "__main__":
    p = HERE / "lecture_sample.pptx"
    d = HERE / "lecture_sample.docx"
    print(f"wrote {p.name} ({build_pptx(p)} bytes)")
    print(f"wrote {d.name} ({build_docx(d)} bytes)")
