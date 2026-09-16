"""PowerPoint parsing via python-pptx."""

import io
import logging

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from app.models import SourceRef
from app.parsers.base import ParsedUnit, ParseResult

logger = logging.getLogger(__name__)


class PptxParser:
    def parse(self, data: bytes) -> ParseResult:
        try:
            presentation = Presentation(io.BytesIO(data))
        except PackageNotFoundError as exc:
            raise ValueError("File is not a readable .pptx presentation") from exc

        result = ParseResult()
        slides = list(presentation.slides)
        result.unit_count = len(slides)

        for index, slide in enumerate(slides):
            slide_no = index + 1
            title = self._title(slide)
            body = self._shape_text(slide, skip=title)
            notes = self._notes(slide)

            # Speaker notes carry the actual explanation far more often than the
            # slide itself, which is usually four bullet fragments. Dropping them
            # would leave the richest text in the deck unsearchable.
            parts = [p for p in (body, notes) if p]
            text = "\n\n".join(parts)

            result.units.append(
                ParsedUnit(
                    text=text,
                    source=SourceRef(kind="slide", slide_no=slide_no),
                    section_title=title,
                    # A slide that is purely an image has no text layer; flag it
                    # so the Phase 4 OCR pass can fill it in.
                    needs_ocr=not text.strip() and self._has_picture(slide),
                )
            )

        return result

    @staticmethod
    def _title(slide) -> str | None:
        try:
            placeholder = slide.shapes.title
        except (AttributeError, ValueError):
            return None
        if placeholder is None or not placeholder.has_text_frame:
            return None
        text = placeholder.text_frame.text.strip()
        return text or None

    @staticmethod
    def _shape_text(slide, skip: str | None) -> str:
        """Collect text from every shape, including inside tables and groups."""
        lines: list[str] = []

        def walk(shapes):
            for shape in shapes:
                # Groups nest arbitrarily; a flat iteration silently loses their content.
                if shape.shape_type == 6 and hasattr(shape, "shapes"):
                    walk(shape.shapes)
                    continue
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        if any(cells):
                            lines.append(" | ".join(cells))
                    continue
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text and text != skip:
                        lines.append(text)

        walk(slide.shapes)
        return "\n".join(lines)

    @staticmethod
    def _notes(slide) -> str:
        if not slide.has_notes_slide:
            return ""
        frame = slide.notes_slide.notes_text_frame
        if frame is None:
            return ""
        text = frame.text.strip()
        return f"Speaker notes: {text}" if text else ""

    @staticmethod
    def _has_picture(slide) -> bool:
        # 13 is MSO_SHAPE_TYPE.PICTURE
        return any(shape.shape_type == 13 for shape in slide.shapes)
